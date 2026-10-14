from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

src="/mnt/data/HOA_Dashboard_QA_Tracker_v4.xlsx"
out="/mnt/data/HOA_Dashboard_QA_Tracker_v5.xlsx"
wb=load_workbook(src)
for n in ["Edit Unit & Ownership QA","Edit Unit Summary"]:
    if n in wb.sheetnames: del wb[n]
ws=wb.create_sheet("Edit Unit & Ownership QA")
headers=["TC ID","Module","Scenario","Preconditions","Test Data","Steps","Expected Result","Actual Result","Status","Priority","Bug ID","Tester","Evidence","Remarks"]
ws.append(headers)

raw = [
("EDIT","Edit Unit","Open Edit Unit","Existing unit","N/A","Open Properties → Edit Unit","Current property, unit, address, bank, owner and balance data load."),
("EDIT","Edit Unit","Save readiness","Existing unit","Valid data","Open Edit Unit","Readiness correctly identifies Property Name, Unit Number, Address and Deposit Account."),
("EDIT","Edit Unit","Save valid changes","Edit permission","Valid change","Modify editable fields → Save","Changes persist after reload."),
("EDIT","Edit Unit","Required field validation","Edit page open","Blank required field","Clear required field → Save","Validation prevents invalid save."),
("EDIT","Edit Unit","Data persistence","Existing unit","Known values","Edit → Save → reopen","Saved values remain accurate."),
("BAL","Starting Balance","Lock when owner assigned","At least one owner","Current owner","Open Edit Unit","Starting Balance and Balance as of Date are locked/read-only."),
("BAL","Starting Balance","Locked fields reject edits","Owner assigned","New amount/date","Attempt to edit locked fields","Fields do not accept changes and no update is persisted."),
("BAL","Starting Balance","Locked explanation","Owner assigned","N/A","Inspect section","UI explains fields are locked because owner/billing is already configured."),
("BAL","Starting Balance","Open when no owner","No owners","N/A","Open Edit Unit","Starting Balance and Balance as of Date are editable/selectable."),
("BAL","Starting Balance","Set balance with no owner","No owners","$500 + valid date","Enter values → Save","Balance is saved successfully."),
("BAL","Starting Balance","Zero balance","No owners","$0 + valid date","Enter values → Save","Zero balance is accepted and persisted."),
("BAL","Starting Balance","Negative balance","No owners","-$30 + valid date","Enter values → Save","Negative balance follows displayed credit/advance rule."),
("BAL","Starting Balance","Date validation","No owners","Blank/invalid date","Enter balance → Save","Validation prevents incomplete starting-balance setup."),
("BAL","Starting Balance","Lock after owner association","No owner initially","Valid owner","Set balance/date → add owner → reload","Starting balance fields become locked after owner association."),
("OWN","Associate Owners","Search by name","Owners exist","Owner name","Search","Matching owners display."),
("OWN","Associate Owners","Search by email","Owners exist","Owner email","Search","Matching owners display."),
("OWN","Associate Owners","Search by property","Owners/units exist","Property name","Search/filter","Matching records display."),
("OWN","Associate Owners","Search by unit","Owners/units exist","Unit name/number","Search/filter","Matching records display."),
("OWN","Associate Owners","No-result search","Owners exist","Unknown text","Search","Correct empty state displays."),
("OWN","Associate Owners","Associated owner cannot be selected","Owner already associated","Current owner","Open list","Selection checkbox/control is disabled/unavailable."),
("OWN","Associate Owners","Unassociated owner can be selected","Eligible owner exists","New owner","Open list","Selection control is enabled."),
("OWN","Associate Owners","Select multiple owners","Multiple unassociated owners","A + B","Select both → Save","Multiple owners are associated without overwriting each other."),
("OWN","Associate Owners","Mixed associated/unassociated","Some associated","Existing + new","Open list","Existing selection disabled; new owner selection enabled."),
("OWN","Associate Owners","Association persistence","Owner selected","Owner A","Save → reload","Owner remains associated."),
("OWN","Associate Owners","Unverified owner assignment","Pending owner exists","Pending owner","Select → Save","Owner may be assigned, but invoice remains blocked until verification."),
("BANK","Deposit Bank Account","Multiple accounts available","2+ bank accounts","A/B","Open section","Available accounts display."),
("BANK","Deposit Bank Account","Change deposit account","2+ accounts","A → B","Select B → Save","Deposit account changes to B."),
("BANK","Deposit Bank Account","Future invoice uses new account","Future invoice exists","A → B","Change account → inspect future invoice","Future billing/payment routing uses new account."),
("BANK","Deposit Bank Account","Due-today invoice","Due-today invoice exists","A → B","Change account → inspect","Current due-today routing follows new account."),
("BANK","Deposit Bank Account","Overdue invoice","Overdue invoice exists","A → B","Change account → inspect","Overdue routing follows new account."),
("BANK","Deposit Bank Account","Partial-paid invoice","Partial-paid invoice exists","A → B","Change account → inspect","Partial-paid routing follows new account."),
("BANK","Deposit Bank Account","Paid invoice historical account","Paid invoice exists before change","A → B","Change → inspect paid invoice","Paid historical invoice remains tied to previous account."),
("BANK","Deposit Bank Account","Bank change activity","Activity timeline enabled","A → B","Change account → inspect timeline","Timeline records old and new account, actor/time."),
("BANK","Deposit Bank Account","Account persistence","Account changed","A → B","Save → reload","B remains selected."),
("ADD","Add Owner","Add owner when owner already exists","At least one owner","New owner","Click Add Owner","Additional owner flow opens."),
("ADD","Add Owner","Search owner by name","Owners exist","Name","Search","Matching owners display."),
("ADD","Add Owner","Search owner by email","Owners exist","Email","Search","Matching owners display."),
("ADD","Add Owner","Select multiple new owners","Multiple eligible owners","A + B","Select → Continue","Multiple selected owners retained."),
("ADD","Add Owner","Effective date cannot be future","Add Owner flow","Future date","Try future date","Future date cannot be selected."),
("ADD","Add Owner","Effective date follows Balance as of Date/current option","Add Owner flow","Balance as of Date + current","Open date selector","Allowed dates follow configured rule; no future date."),
("ADD","Add Owner","Existing owner remains active","Existing owner","A + B","Add B → complete","A remains active; B is added."),
("ADD","Add Owner","Historical invoices remain with prior owner","Historical invoices exist","A history + B","Add B → inspect","Past invoices/payments stay with owners active at the time."),
("ADD","Add Owner","Outstanding balance shared forward","Outstanding balance exists","$430 + B","Add B → inspect","Selected owners share current outstanding balance from effective date forward."),
("ADD","Add Owner","Impact Summary","Existing + new owner","Known balance","Review summary","Summary accurately describes historical and forward balance treatment."),
("CHG","Change Ownership","Open transfer","Current owner exists","Current owner","Click Change Ownership","Current owners shown as releasing owners."),
("CHG","Change Ownership","Search by name","Owners exist","Name","Search","Matching owners display."),
("CHG","Change Ownership","Search by email","Owners exist","Email","Search","Matching owners display."),
("CHG","Change Ownership","Search by phone","Owners exist","Phone","Search","Matching owners display."),
("CHG","Change Ownership","Multiple new owners","Eligible owners exist","A+B","Select both","Multiple new owners can be selected."),
("CHG","Change Ownership","Transfer date","Current owner","Valid/invalid dates","Set transfer date","Date validation follows ownership timeline rules."),
("CHG","Change Ownership","Migrate unit balance","$430 outstanding","$430","Select Migrate unit balance","Balance remains; unpaid amounts carry forward."),
("CHG","Change Ownership","Reset unit balance","$430 outstanding","$430","Select Reset unit balance","Due/overdue balances are settled as specified; credits handled as specified."),
("CHG","Change Ownership","Reset requires payment method","Outstanding balance","No method","Select Reset → Continue","Payment method validation appears."),
("CHG","Change Ownership","Settlement method","Outstanding balance","Cash/check/money order/other","Select method → confirm","Settlement is recorded with selected method."),
("CHG","Change Ownership","Keep separate active charges","Outstanding/partial charges","Partial invoice","Select option","Outstanding charges migrate; partial charges re-created with current balance."),
("CHG","Change Ownership","Keep future charges","Future charges exist","Scheduled invoice","Select option","Future charges remain and apply after transfer where applicable."),
("CHG","Change Ownership","Keep recurring templates","Recurring templates exist","Monthly template","Select option","Recurring templates continue."),
("CHG","Change Ownership","Historical portal visibility","Transfer completed","Effective date","Login as new owner","New owner cannot view periods before effective date."),
("CHG","Change Ownership","Old owner released","Transfer completed","Effective date","Inspect old owner","Old owner inactive after effective date."),
("CHG","Change Ownership","Migration impact summary","Transfer setup","Migrate + keep options","Review summary","Summary matches selected options."),
("CHG","Change Ownership","Reset impact summary","Transfer setup","Reset + settlement","Review summary","Summary matches settlement and transfer options."),
("CHG","Change Ownership","History preserved","Transfer completed","Historical invoices/payments","Inspect history","Historical records remain preserved."),
("REM","Remove Owner","Only-owner warning","Exactly one owner","Owner A","Click Remove Owner","Warning identifies only-owner condition."),
("REM","Remove Owner","Cancel future scheduled invoices","Only owner + future invoices","Future invoices","Confirm removal","Future scheduled invoices cancelled; generated invoices unchanged."),
("REM","Remove Owner","History preserved","Historical invoices exist","Past invoices","Remove → inspect","Invoice history preserved."),
("REM","Remove Owner","Outstanding balance displayed","$430 outstanding","$430","Open removal","Correct balance shown."),
("REM","Remove Owner","Keep balance","Outstanding due/overdue","$430","Select Keep balance → confirm","Due/overdue invoices remain unchanged."),
("REM","Remove Owner","Mark as paid","Outstanding due/overdue","$430 + method","Select Mark paid → method → confirm","Invoices marked paid and settlement recorded."),
("REM","Remove Owner","Waive balance","Outstanding balance","$430","Select Waive → confirm","Amount waived; history/waiver details preserved."),
("REM","Remove Owner","Payment method required","Outstanding balance","Blank method","Mark paid → continue","Validation requires method."),
("REM","Remove Owner","Payment method options","Outstanding balance","Cash/check/money order/other","Open selector","Expected methods available."),
("REM","Remove Owner","Past removal date blocked","Owner active","Past date","Try past date","Past date not allowed."),
("REM","Remove Owner","Today removal date","Owner active","Today","Select today","Today accepted."),
("REM","Remove Owner","Future removal date","Owner active","Future date","Select future date","Future date accepted."),
("REM","Remove Owner","Remove one of multiple owners","2+ owners","Remove A","Complete removal","A becomes inactive; remaining owners stay active."),
("REM","Remove Owner","Open unpaid invoices reassigned","2+ owners + unpaid invoice","Invoice","Remove A","Open unpaid invoices move to remaining owners."),
("REM","Remove Owner","Paid invoices retain history","Paid invoice exists","Invoice","Remove A → inspect","Paid invoice retains historical owner."),
("REM","Remove Owner","In-progress invoices retain history","Processing invoice exists","Invoice","Remove A → inspect","In-progress invoice retains historical owner."),
("REM","Remove Owner","Impact Summary","Removal setup","Handling option","Review summary","Summary accurately states invoice reassignment and ownership end."),
("REM","Remove Owner","Owner notification","Owner has email","Removal","Complete removal","Owner receives email notification."),
("REM","Remove Owner","Last owner becomes ownerless","One owner","Owner A","Remove A","Unit becomes ownerless after effective date."),
("REM","Remove Owner","Cancel removal","Removal modal open","N/A","Cancel/close","No ownership or invoice changes."),
("CREDIT","Remove Owner - Credit","Detect prepaid credit","Only owner + $30 credit","$30","Open removal","Credit is displayed."),
("CREDIT","Remove Owner - Credit","Post credit as HOA expense","$30 credit","$30","Confirm removal","Credit posted as pending HOA expense and cleared."),
("CREDIT","Remove Owner - Credit","Credit + outstanding handling","$30 credit + $430 due","Known balances","Complete removal","Both credit-clearing and selected balance treatment apply correctly."),
("CREDIT","Remove Owner - Credit","Credit audit","Credit exists","$30","Complete → inspect activity","Credit expense/clearing action recorded."),
("AUDIT","Ownership Audit","Add Owner activity","Add Owner completed","A → B","Inspect timeline","Timeline records owner addition/effective date/actor."),
("AUDIT","Ownership Audit","Change Ownership activity","Transfer completed","A → B","Inspect timeline","Timeline records ownership change/effective date."),
("AUDIT","Ownership Audit","Remove Owner activity","Removal completed","A removed","Inspect timeline","Timeline records removal/effective date/handling."),
("AUDIT","Ownership Audit","Bank account activity","A → B","A → B","Inspect timeline","Timeline records bank change from A to B."),
("AUDIT","Ownership Audit","No duplicate owner after retry","Owner selected","Same owner","Save/retry/reload","No duplicate association."),
("AUDIT","Ownership Audit","No duplicate invoice after ownership action","Billing configured","Repeated action","Repeat/refresh/retry","No unintended duplicate invoice."),
("AUDIT","Ownership Audit","Historical ownership preserved","Multiple periods","Historical data","Inspect history","Historical ownership/billing remains intact."),
("AUDIT","Ownership Audit","Ownership actions hidden when ownerless","No owner","N/A","Open Edit Unit","Add Owner, Change Ownership, Remove Owner are not visible."),
("AUDIT","Ownership Audit","Ownership actions visible with owner","At least one owner","N/A","Open Edit Unit","Ownership action controls are visible."),
("AUDIT","Ownership Audit","Permission check","No edit permission","N/A","Open Edit Unit","Unauthorized ownership/bank/balance changes are prevented.")
]
for i,(prefix,module,scenario,pre,data,steps,exp) in enumerate(raw,1):
    ws.append([f"{prefix}-{i:03d}",module,scenario,pre,data,steps,exp,"","Not Tested","P0","","","",""])

navy="1F4E78"; white="FFFFFF"; thin=Side(style="thin",color="B7B7B7")
for c in ws[1]:
    c.fill=PatternFill("solid",fgColor=navy); c.font=Font(color=white,bold=True)
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:N{ws.max_row}"; ws.sheet_view.showGridLines=False
for i,w in enumerate([18,24,40,42,34,48,62,42,16,12,14,18,24,34],1):
    ws.column_dimensions[get_column_letter(i)].width=w

dv=DataValidation(type="list",formula1='"Not Tested,In Progress,Passed,Failed,Blocked,Not Applicable,Retest Required"'); ws.add_data_validation(dv); dv.add(f"I2:I{ws.max_row}")
dp=DataValidation(type="list",formula1='"P0,P1,P2"'); ws.add_data_validation(dp); dp.add(f"J2:J{ws.max_row}")
for s,c in {"Passed":"E2F0D9","Failed":"F4CCCC","Blocked":"FCE4D6","In Progress":"FFF2CC","Not Tested":"E7E6E6","Retest Required":"FCE4D6"}.items():
    ws.conditional_formatting.add(f"I2:I{ws.max_row}",FormulaRule(formula=[f'$I2="{s}"'],fill=PatternFill("solid",fgColor=c)))

sm=wb.create_sheet("Edit Unit Summary"); sm.sheet_view.showGridLines=False
sm.merge_cells("A1:D1"); sm["A1"]="Edit Unit & Ownership QA Summary"; sm["A1"].font=Font(size=18,bold=True)
items=[
("Total Test Cases",f"=COUNTA('Edit Unit & Ownership QA'!A2:A{ws.max_row})"),
("Passed",f'=COUNTIF(\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Passed")'),
("Failed",f'=COUNTIF(\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Failed")'),
("Blocked",f'=COUNTIF(\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Blocked")'),
("Not Tested",f'=COUNTIF(\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Not Tested")'),
("Coverage %","=IF(B2=0,0,(B3+B4+B5)/B2)"),
("Pass Rate %","=IF((B3+B4+B5)=0,0,B3/(B3+B4+B5))"),
("P0 Failed/Blocked",f'=COUNTIFS(\'Edit Unit & Ownership QA\'!J2:J{ws.max_row},"P0",\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Failed")+COUNTIFS(\'Edit Unit & Ownership QA\'!J2:J{ws.max_row},"P0",\'Edit Unit & Ownership QA\'!I2:I{ws.max_row},"Blocked")')]
for r,(label,formula) in enumerate(items,2):
    sm.cell(r,1,label).font=Font(bold=True); sm.cell(r,2,formula)
sm["B7"].number_format="0.00%"; sm["B8"].number_format="0.00%"
sm.column_dimensions["A"].width=30; sm.column_dimensions["B"].width=20
wb.save(out)
print(f"Created {out} with {len(raw)} Edit Unit & Ownership test cases.")
