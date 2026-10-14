from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from pathlib import Path

path = Path("/mnt/data/HOA_Dashboard_QA_Tracker.xlsx")
wb = load_workbook(path)

if "Manage HOA Fee QA" in wb.sheetnames:
    del wb["Manage HOA Fee QA"]

ws = wb.create_sheet("Manage HOA Fee QA")

headers = [
    "TC ID","Module","Scenario","Preconditions","Test Data","Steps",
    "Expected Result","Actual Result","Status","Priority","Bug ID",
    "Tester","Evidence","Remarks"
]

cases = []

def add(tc, scenario, pre, data, steps, expected, priority="High"):
    cases.append([
        tc, "Manage HOA Fee", scenario, pre, data, steps, expected,
        "", "Not Started", priority, "", "", "", ""
    ])

# Configuration access / page
add("FEE-001","Open Manage HOA Fee",
    "User is logged in as an authorized HOA Board Member",
    "Settings → Manage HOA Fee",
    "Open Settings → click Manage HOA Fee",
    "Manage HOA Fee configuration page opens successfully.")
add("FEE-002","Verify configuration summary",
    "Manage HOA Fee page is open",
    "Recurring fee, invoices/year, due day, partial payments, late fees",
    "Review Configuration Summary",
    "Summary reflects the saved configuration values.")
add("FEE-003","Verify recurring fee section",
    "Manage HOA Fee page is open",
    "$100/month, Monthly, 1 Year, 13th of month",
    "Review recurring fee controls",
    "Fee amount, billing frequency, tenure, and due day controls are displayed and reflect saved values.")
add("FEE-004","Verify recurring fee enable/disable",
    "Manage HOA Fee page is open",
    "Recurring fee toggle",
    "Turn recurring fee off/on",
    "Toggle state is saved/handled correctly and invoice generation behavior follows the configured state.")
add("FEE-005","Verify partial payments setting",
    "Manage HOA Fee page is open",
    "Accept Partial Payments",
    "Toggle partial payments",
    "Setting can be changed and saved; invoice/payment behavior follows the saved setting.")

# Recurring fee
add("FEE-006","Validate recurring fee amount",
    "Recurring fee section is enabled",
    "Positive fee amount",
    "Enter valid fee amount → save",
    "Valid amount is accepted and reflected in the configuration summary.")
add("FEE-007","Reject invalid recurring fee amount",
    "Recurring fee section is enabled",
    "Negative/non-numeric amount",
    "Enter invalid amount → save",
    "Invalid amount is rejected with clear validation.")
add("FEE-008","Verify billing frequency",
    "Recurring fee section is enabled",
    "Monthly",
    "Review/select billing frequency",
    "Selected billing frequency is stored and reflected in generated invoice schedule.")
add("FEE-009","Verify tenure",
    "Recurring fee section is enabled",
    "1 Year",
    "Review/select tenure",
    "Selected tenure is stored and used for invoice generation according to the configuration.")
add("FEE-010","Verify due day of month",
    "Recurring fee section is enabled",
    "13th of every month",
    "Select due day → save",
    "Due day is stored and generated invoices use the configured due day.")
add("FEE-011","Verify invalid day handling",
    "Recurring fee section is enabled",
    "Day values including 31",
    "Configure a month-end-sensitive due day",
    "System handles months without that calendar day according to the defined business rule and does not create invalid dates.")
add("FEE-012","Verify configuration save",
    "User has permission to edit settings",
    "Changed fee configuration",
    "Modify fee settings → save/confirm",
    "Configuration is persisted and remains after refresh/re-login.")
add("FEE-013","Verify configuration after refresh",
    "Configuration was saved",
    "Saved values",
    "Refresh page",
    "Saved values remain unchanged and summary matches the stored configuration.")

# Late fees
add("FEE-014","Verify Late Fee Settings enabled state",
    "Manage HOA Fee page is open",
    "Late fee toggle",
    "Turn late fees on/off",
    "Late fee settings state is correctly reflected and saved.")
add("FEE-015","Verify percentage late fee",
    "Late fees enabled",
    "10% of Total Amount",
    "Configure percentage late fee = 10%",
    "Late fee rule is saved with the configured percentage.")
add("FEE-016","Verify grace period",
    "Late fees enabled",
    "0 days / positive days",
    "Configure grace period",
    "Late fee is applied only after the configured grace period.")
add("FEE-017","Verify recurring late fee toggle",
    "Late fees enabled",
    "Repeat Yes/No",
    "Enable/disable repeat",
    "Repeat behavior is stored and late fee generation follows the setting.")
add("FEE-018","Verify late fee end day",
    "Late fee rule supports end day",
    "No end / configured end day",
    "Configure end day → save",
    "Late fees stop according to the configured end-day rule.")
add("FEE-019","Verify late fee start day",
    "Late fee rule supports start day",
    "14th",
    "Configure start day = 14th",
    "Late fee is eligible from the configured start day.")
add("FEE-020","Verify applicable properties/units - all",
    "Late fee rule is enabled",
    "All Properties",
    "Select All Properties → save",
    "Rule applies to all eligible properties/units.")
add("FEE-021","Verify applicable properties/units - specific",
    "Late fee rule is enabled and properties/units exist",
    "Specific properties/units",
    "Select Specific Properties → choose records → save",
    "Rule applies only to selected properties/units.")
add("FEE-022","Verify multiple late fee rules",
    "Late fee settings enabled",
    "Fee 1, Fee 2",
    "Add another late fee rule → configure → save",
    "Multiple rules are stored and evaluated according to the configured order/business rules.")
add("FEE-023","Delete late fee rule",
    "At least two late fee rules exist",
    "Fee rule",
    "Delete one rule → save",
    "Selected rule is removed and remaining rules stay intact.")
add("FEE-024","Verify late fee calculation example",
    "Fee = $100, late fee = 10%, no grace period, start day = 14th",
    "Invoice $100 due 13th; unpaid on 14th",
    "Create/use qualifying invoice → pass the configured late-fee trigger",
    "Late fee is $10 and is applied according to the configured rule.")
add("FEE-025","Verify late fee does not apply during grace period",
    "Grace period > 0",
    "Grace = 5 days",
    "Use unpaid invoice within first 5 days after due date",
    "Late fee is not applied before the grace period expires.")
add("FEE-026","Verify late fee after grace period",
    "Grace period > 0",
    "Grace = 5 days",
    "Use unpaid invoice after grace period",
    "Late fee is applied at the correct time.")
add("FEE-027","Verify repeated late fee behavior",
    "Repeat enabled",
    "Repeat = Yes; configured frequency",
    "Leave invoice unpaid across multiple configured cycles",
    "Additional late fees occur according to the configured repeat frequency and stop/end rules.")
add("FEE-028","Verify no repeated late fee when Repeat is disabled",
    "Repeat disabled",
    "Repeat = No",
    "Leave invoice unpaid after first late fee",
    "No additional recurring late fee is generated.")
add("FEE-029","Verify late fee calculation uses outstanding amount",
    "Late fee rule is percentage-based and invoice is partially paid",
    "Invoice $100, payment $40, late fee 10%",
    "Partially pay invoice → reach late-fee trigger",
    "Late fee calculation follows the product's defined base (e.g., outstanding amount if configured/expected) and does not incorrectly charge on already-paid amount.")
add("FEE-030","Verify late fee applies only to eligible properties/units",
    "Specific property/unit rule is configured",
    "Eligible + ineligible units",
    "Generate overdue invoices for both → trigger late fee",
    "Late fee applies only to eligible properties/units.")

# Payment reminders
add("FEE-031","Verify Payment Reminders enabled",
    "Manage HOA Fee page is open",
    "Payment Reminders toggle",
    "Enable/disable Payment Reminders",
    "Reminder setting is saved and controls reminder generation.")
add("FEE-032","Verify notification channels",
    "Payment Reminders enabled",
    "Email, In-app",
    "Toggle Email and In-app independently",
    "Selected channels are saved and reminders are delivered through only the enabled channels.")
add("FEE-033","Verify before-due-date reminder",
    "Payment Reminders enabled",
    "5 days before due date",
    "Enable 5-day reminder → use qualifying invoice",
    "Reminder is generated at the configured lead time.")
add("FEE-034","Verify multiple before-due reminders",
    "Payment Reminders enabled",
    "5, 3, 1 days before due date",
    "Enable multiple lead times",
    "Each enabled reminder is generated at its configured time without generating disabled reminders.")
add("FEE-035","Verify on-due-date reminder",
    "Payment Reminders enabled",
    "On the due date",
    "Enable on-due-date reminder",
    "Reminder is generated on the configured due date.")
add("FEE-036","Verify after-due-date reminder",
    "Payment Reminders enabled",
    "1 day and 7 days after due date",
    "Enable overdue reminders",
    "Reminders are generated at the configured overdue intervals for unpaid invoices.")
add("FEE-037","Verify continue overdue reminders",
    "Continue overdue reminders enabled",
    "Scheduled continuation",
    "Leave invoice unpaid over multiple cycles",
    "Reminder continues according to the configured schedule until the configured stopping condition.")
add("FEE-038","Verify overdue frequency",
    "Continue overdue reminders enabled",
    "7 days",
    "Set overdue frequency = 7 days",
    "Recurring overdue reminders follow the 7-day interval.")
add("FEE-039","Verify maximum overdue reminders",
    "Continue overdue reminders enabled",
    "Max overdue reminders = 1",
    "Leave invoice unpaid",
    "No more than the configured maximum number of overdue reminders is generated.")
add("FEE-040","Verify grace period before overdue reminders",
    "Grace period consideration is enabled",
    "0 days or configured grace",
    "Leave invoice unpaid during grace period",
    "Overdue reminders do not trigger before the configured grace period has elapsed.")
add("FEE-041","Verify paid invoice stops reminders",
    "Reminder workflow active",
    "Invoice paid before next reminder",
    "Pay invoice → wait/trigger next reminder cycle",
    "No further unpaid-invoice reminder is generated for the paid invoice.")
add("FEE-042","Verify partially paid invoice reminder behavior",
    "Partial payments enabled and reminder workflow active",
    "Invoice partially paid",
    "Partially pay invoice → reach reminder trigger",
    "Reminder behavior follows the configured partial-payment policy; system does not treat the invoice incorrectly as fully paid.")
add("FEE-043","Verify custom reminder subject",
    "Payment Reminders enabled",
    "Custom subject",
    "Enter custom subject → save → trigger reminder",
    "Notification uses the saved custom subject.")
add("FEE-044","Verify custom reminder body",
    "Payment Reminders enabled",
    "Custom body",
    "Enter custom body → save → trigger reminder",
    "Notification uses the saved custom body.")
add("FEE-045","Verify invoice placeholder in custom subject",
    "Custom reminder subject supports invoiceNumber placeholder",
    "Use {{invoiceNumber}} or product-supported syntax",
    "Save custom subject → trigger reminder",
    "Supported invoice number placeholder is replaced with the correct invoice number.")
add("FEE-046","Verify default message when custom body is blank",
    "Custom body is optional",
    "Blank custom body",
    "Leave custom body blank → trigger reminder",
    "Default reminder message is used.")
add("FEE-047","Verify separate templates",
    "Payment Reminders enabled",
    "Before due / Due date / Overdue templates",
    "Configure different subject/body values for each stage → trigger each",
    "Correct template is used for each reminder stage.")

# Critical configuration completion / confirmation
add("FEE-048","Verify Manage HOA Fee configuration completion state",
    "User can configure HOA Fee settings",
    "Valid recurring fee + billing frequency + tenure + due day and required confirmation/save",
    "Complete required Manage HOA Fee configuration → save/confirm → return to Settings/Dashboard",
    "System records the Manage HOA Fee setup step as completed and exposes the correct completed state.","Critical")
add("FEE-049","Incomplete Manage HOA Fee configuration blocks property creation",
    "Manage HOA Fee setup is not completed/confirmed",
    "Incomplete HOA fee configuration",
    "Attempt to create/add a property",
    "Property creation is blocked if Manage HOA Fee completion is a required prerequisite; user receives a clear message explaining that Manage HOA Fee setup must be completed first.","Critical")
add("FEE-050","Incomplete Manage HOA Fee configuration blocks invoice creation",
    "Manage HOA Fee setup is not completed/confirmed",
    "Incomplete HOA fee configuration",
    "Attempt to create an invoice",
    "Invoice creation is blocked if Manage HOA Fee completion is required; user is directed to complete Manage HOA Fee configuration first.","Critical")
add("FEE-051","Incomplete Manage HOA Fee configuration blocks recurring invoice generation",
    "Manage HOA Fee setup is incomplete",
    "No completed billing configuration",
    "Trigger scheduled/automatic recurring invoice generation",
    "No recurring invoice is generated from an incomplete configuration; system records/handles the blocked state safely.","Critical")
add("FEE-052","Incomplete Manage HOA Fee configuration blocks owner assignment when assignment triggers billing",
    "Manage HOA Fee setup is incomplete and product requires completed fee configuration before billing-linked owner/unit assignment",
    "New owner + unit",
    "Create owner → attempt to assign unit",
    "Assignment is blocked only if the defined prerequisite is required for billing-linked ownership; message clearly identifies the missing Manage HOA Fee configuration.","High")
add("FEE-053","Incomplete Manage HOA Fee configuration blocks billing-related actions",
    "Manage HOA Fee setup is incomplete",
    "Property/invoice/payment-related action",
    "Attempt billing-related actions exposed by the application",
    "Only actions that depend on a completed HOA fee configuration are blocked; unrelated administrative actions remain available.")
add("FEE-054","Complete Manage HOA Fee unlocks property creation",
    "Manage HOA Fee configuration is fully completed/confirmed",
    "Valid fee configuration",
    "Complete/confirm Manage HOA Fee → open Properties → Add Property",
    "Property creation becomes available without the prerequisite block.","Critical")
add("FEE-055","Complete Manage HOA Fee unlocks invoice creation",
    "Manage HOA Fee configuration is fully completed/confirmed",
    "Valid fee configuration",
    "Complete/confirm Manage HOA Fee → open Invoice creation",
    "Invoice creation becomes available and uses the configured billing rules.","Critical")
add("FEE-056","Configuration changes after completion",
    "Manage HOA Fee was previously completed",
    "Change fee amount/due day/late fee",
    "Edit configuration → save → review completion state",
    "Configuration remains in a valid completed state if required fields remain valid; downstream behavior uses the updated settings.")
add("FEE-057","Invalid configuration cannot be marked complete",
    "Manage HOA Fee page is open",
    "Missing required billing configuration",
    "Attempt to save/confirm incomplete configuration",
    "System does not mark the setup step complete and displays required validation.")
add("FEE-058","Configuration completion persists after logout/login",
    "Manage HOA Fee setup is completed",
    "Completed configuration",
    "Logout → login → open Dashboard/Settings",
    "Completion state remains correct after re-authentication.")
add("FEE-059","Configuration completion is HOA-specific",
    "Two HOAs exist with separate settings",
    "HOA A complete; HOA B incomplete",
    "Switch between HOA A and HOA B → attempt property/invoice actions",
    "Completion state and access gating are evaluated independently per HOA.","Critical")
add("FEE-060","Gating cannot be bypassed by direct URL/API",
    "Manage HOA Fee setup is incomplete",
    "Direct property/invoice route or authorized API request",
    "Attempt to open/create property or invoice using direct navigation/API",
    "Server-side authorization/business-rule validation enforces the prerequisite; the restriction is not only a frontend/UI check.","Critical")
add("FEE-061","Gating message redirects to Manage HOA Fee",
    "Manage HOA Fee setup is incomplete",
    "Blocked property/invoice action",
    "Attempt blocked action → select provided setup link/action",
    "User is redirected to Manage HOA Fee configuration.")
add("FEE-062","Dashboard setup progress reflects Manage HOA Fee completion",
    "Dashboard setup checklist is enabled",
    "Manage HOA Fee incomplete then complete",
    "Check dashboard before and after completing Manage HOA Fee",
    "Dashboard setup progress and remaining-step indicator update correctly.")
add("FEE-063","Owner creation + assignment respects Manage HOA Fee prerequisite",
    "Manage HOA Fee setup incomplete",
    "New owner + selected unit",
    "Create owner and assign unit",
    "System follows the defined prerequisite: if completed fee configuration is required before billing-linked assignment, assignment/billing is blocked until setup is complete; no unintended invoice is generated.","Critical")
add("FEE-064","Owner assignment after completion generates invoice according to settings",
    "Manage HOA Fee setup is completed and recurring fee is active",
    "Owner + unit + $100 monthly fee + 13th due day",
    "Create/assign owner to unit → inspect invoices",
    "Invoice is generated according to the saved fee amount, billing frequency, tenure, due day and applicable late/reminder rules.","Critical")

# Settings navigation
add("FEE-065","Settings card opens Manage HOA Fee",
    "Settings page is accessible",
    "Manage HOA Fee card",
    "Click Manage HOA Fee",
    "Correct Manage HOA Fee page opens.")
add("FEE-066","Back to Settings navigation",
    "Manage HOA Fee page is open",
    "Back to Settings",
    "Click Back to Settings",
    "Settings page opens and displays the Manage HOA Fee status.")
add("FEE-067","Settings summary reflects configuration completion",
    "Manage HOA Fee configuration is completed",
    "Completed state",
    "Open Settings",
    "Manage HOA Fee card/status reflects the saved configuration state and does not incorrectly show it as incomplete.")

# End-to-end
add("FEE-068","End-to-end prerequisite gating workflow",
    "New HOA exists with Manage HOA Fee incomplete",
    "Incomplete → Complete → Property → Owner → Invoice",
    "1. Attempt Add Property → 2. Verify blocked → 3. Attempt Create Invoice → 4. Verify blocked → 5. Complete Manage HOA Fee → 6. Verify setup completed → 7. Add Property → 8. Create/assign Owner → 9. Verify invoice generation",
    "Before completion, dependent property/invoice/billing actions are blocked as defined. After completion, dependent workflows are unlocked and use the configured HOA fee settings.","Critical")

# Write
for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c, h)
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin = Side(style="thin", color="D9E1F2")
for r, row in enumerate(cases, 2):
    for c, value in enumerate(row, 1):
        cell = ws.cell(r, c, value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[r].height = 58

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:N{ws.max_row}"

widths = {
    "A":12, "B":22, "C":44, "D":38, "E":34, "F":62,
    "G":76, "H":42, "I":16, "J":14, "K":14, "L":18, "M":26, "N":34
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

status_dv = DataValidation(type="list", formula1='"Not Started,Pass,Fail,Blocked,Not Applicable,Retest"', allow_blank=False)
priority_dv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=False)
ws.add_data_validation(status_dv)
ws.add_data_validation(priority_dv)
status_dv.add(f"I2:I{ws.max_row}")
priority_dv.add(f"J2:J{ws.max_row}")

ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Pass"'], fill=PatternFill("solid", fgColor="C6EFCE"))
)
ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Fail"'], fill=PatternFill("solid", fgColor="FFC7CE"))
)
ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Blocked"'], fill=PatternFill("solid", fgColor="FFEB9C"))
)
ws.conditional_formatting.add(
    f"J2:J{ws.max_row}",
    FormulaRule(formula=['J2="Critical"'], fill=PatternFill("solid", fgColor="F4CCCC"))
)
ws.sheet_view.showGridLines = False

# Summary sheet
if "Manage HOA Fee Summary" in wb.sheetnames:
    del wb["Manage HOA Fee Summary"]
sm = wb.create_sheet("Manage HOA Fee Summary")
sm["A1"] = "Manage HOA Fee - QA Summary"
sm["A1"].font = Font(size=18, bold=True)
sm.merge_cells("A1:D1")

summary = [
    ("Total Test Cases", f"=COUNTA('Manage HOA Fee QA'!A2:A{ws.max_row})"),
    ("Passed", f'=COUNTIF(\'Manage HOA Fee QA\'!I2:I{ws.max_row},"Pass")'),
    ("Failed", f'=COUNTIF(\'Manage HOA Fee QA\'!I2:I{ws.max_row},"Fail")'),
    ("Blocked", f'=COUNTIF(\'Manage HOA Fee QA\'!I2:I{ws.max_row},"Blocked")'),
    ("Not Started", f'=COUNTIF(\'Manage HOA Fee QA\'!I2:I{ws.max_row},"Not Started")'),
    ("Pass Rate", '=IF(B2=0,0,B3/B2)'),
    ("Execution Rate", '=IF(B2=0,0,(B3+B4+B5)/B2)'),
    ("Critical Cases", f'=COUNTIF(\'Manage HOA Fee QA\'!J2:J{ws.max_row},"Critical")'),
    ("Critical Failures", f'=COUNTIFS(\'Manage HOA Fee QA\'!J2:J{ws.max_row},"Critical",\'Manage HOA Fee QA\'!I2:I{ws.max_row},"Fail")'),
]
for i, (label, formula) in enumerate(summary, 2):
    sm.cell(i,1,label).font = Font(bold=True)
    sm.cell(i,2,formula)
    sm.cell(i,1).fill = PatternFill("solid", fgColor="D9EAF7")
    sm.cell(i,1).border = Border(bottom=thin)
    sm.cell(i,2).border = Border(bottom=thin)
sm["B7"].number_format = "0.0%"
sm["B8"].number_format = "0.0%"
sm.column_dimensions["A"].width = 32
sm.column_dimensions["B"].width = 18
sm.sheet_view.showGridLines = False

# Add reference notes
if "Test Reference" in wb.sheetnames:
    ref = wb["Test Reference"]
else:
    ref = wb.create_sheet("Test Reference")

r = ref.max_row + 2
notes = [
    ("Manage HOA Fee prerequisite",
     "For the requested workflow, verify that incomplete/unconfirmed Manage HOA Fee configuration blocks dependent property, invoice, and billing workflows where the product defines this prerequisite."),
    ("Server-side enforcement",
     "The prerequisite should be enforced server-side/API-side, not only by disabling or hiding frontend controls."),
    ("Completion",
     "Completion should be recorded only after required configuration is valid and saved/confirmed."),
    ("Downstream behavior",
     "After completion, property/invoice/owner-assignment workflows should use the saved HOA fee configuration."),
    ("Owner lifecycle",
     "For owner assignment flows, verify invoice creation follows Manage HOA Fee settings and invitation/account status follows the configured lifecycle.")
]
for label, note in notes:
    ref.cell(r,1,label).font = Font(bold=True)
    ref.cell(r,1).fill = PatternFill("solid", fgColor="D9EAF7")
    ref.cell(r,2,note).alignment = Alignment(wrap_text=True)
    r += 1

# Update main summary with this module
dash = wb["Summary Dashboard"]
row = dash.max_row + 2
dash.cell(row,1,"Manage HOA Fee QA")
dash.cell(row,1).font = Font(bold=True)
dash.cell(row,2,f"='Manage HOA Fee Summary'!B2")
dash.cell(row,3,"test cases added")
dash.cell(row,1).fill = PatternFill("solid", fgColor="D9EAF7")

wb.save(path)

print(f"Updated: {path}")
print(f"Manage HOA Fee test cases added: {len(cases)}")
