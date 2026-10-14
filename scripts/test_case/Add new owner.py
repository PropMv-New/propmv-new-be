from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from pathlib import Path

path = Path("/mnt/data/HOA_Dashboard_QA_Tracker.xlsx")
wb = load_workbook(path)

# Replace/add dedicated Add New Owner QA sheet
if "Add Owner QA" in wb.sheetnames:
    del wb["Add Owner QA"]

ws = wb.create_sheet("Add Owner QA")

headers = [
    "TC ID","Module","Scenario","Preconditions","Test Data","Steps",
    "Expected Result","Actual Result","Status","Priority","Bug ID",
    "Tester","Evidence","Remarks"
]

cases = []

def add(tc, scenario, pre, data, steps, expected, priority="High"):
    cases.append([
        tc, "Add New Owner", scenario, pre, data, steps, expected,
        "", "Not Started", priority, "", "", "", ""
    ])

# -------------------------
# Owner Information
# -------------------------
add("ADD-001","Open Add New Owner page",
    "User is logged in as HOA Board Member with permission to create owners",
    "Add Owner",
    "Property Owners → Add Owner",
    "Open Property Owners → Click Add Owner",
    "Add new owner page opens successfully with Owner information, Emergency contact, Unit ownership, Cancel and Save owner sections.")
add("ADD-002","Verify required fields",
    "Add New Owner page is open",
    "First Name, Last Name, Email, Confirm Email, property assignment",
    "Review the form",
    "Required fields are clearly indicated and validation is triggered when required information is missing.")
add("ADD-003","Create owner with valid mandatory information and no property",
    "Add New Owner page is open",
    "Unique first/last name + unique email; no property",
    "Enter valid required fields → Select No for property assignment → Save owner",
    "Owner is created successfully without a property assignment, if the product permits owner creation without a property.")
add("ADD-004","First Name required",
    "Add New Owner page is open",
    "Blank First Name",
    "Leave First Name blank → fill other required fields → Save",
    "Validation message is displayed and owner is not created.")
add("ADD-005","Last Name required",
    "Add New Owner page is open",
    "Blank Last Name",
    "Leave Last Name blank → fill other required fields → Save",
    "Validation message is displayed and owner is not created.")
add("ADD-006","Email required",
    "Add New Owner page is open",
    "Blank Email",
    "Leave Email blank → fill other required fields → Save",
    "Validation message is displayed and owner is not created.")
add("ADD-007","Confirm Email required",
    "Add New Owner page is open",
    "Blank Confirm Email",
    "Leave Confirm Email blank → fill other required fields → Save",
    "Validation message is displayed and owner is not created.")
add("ADD-008","Validate email format",
    "Add New Owner page is open",
    "invalid-email",
    "Enter invalid email → attempt Save",
    "Invalid email format is rejected with a clear validation message.")
add("ADD-009","Confirm email must match",
    "Add New Owner page is open",
    "Email: owner1@example.com; Confirm: owner2@example.com",
    "Enter different Email and Confirm Email → Save",
    "Validation indicates that email and confirmation do not match; owner is not created.")
add("ADD-010","Email uniqueness in database",
    "An owner already exists with the email",
    "Existing owner email",
    "Enter existing email → complete valid form → Save",
    "Owner creation is rejected with a clear duplicate/unique email message; no duplicate owner is created.","Critical")
add("ADD-011","Email uniqueness is case handled correctly",
    "An owner exists with an email",
    "Existing email with different casing",
    "Enter same email using different casing → Save",
    "System applies the intended email uniqueness rule and does not create an unintended duplicate.","Critical")
add("ADD-012","Email with leading/trailing spaces",
    "Add New Owner page is open",
    "Email with spaces",
    "Enter email with leading/trailing spaces → Save",
    "Email is normalized/validated according to the product rule and does not create malformed data.")
add("ADD-013","Valid phone number",
    "Add New Owner page is open",
    "Valid US/international phone",
    "Enter valid phone → Save",
    "Phone is accepted and stored in the expected format.")
add("ADD-014","Phone duplicate validation",
    "Another owner already has the phone number, if phone uniqueness is a business rule",
    "Existing phone number",
    "Enter an already-used phone number → complete form → Save",
    "System applies the configured phone-duplicate rule. If phone numbers are required to be unique, creation is blocked with a clear duplicate message; otherwise the duplicate is allowed.","Critical")
add("ADD-015","Phone optional",
    "Add New Owner page is open",
    "Blank phone",
    "Leave Phone blank → complete required fields → Save",
    "Owner can be created without a phone because Phone is optional.")
add("ADD-016","International phone format",
    "Add New Owner page is open",
    "International phone",
    "Enter international phone number → Save",
    "International format is accepted as indicated by the form.")
add("ADD-017","Verify phone field does not accept invalid input",
    "Add New Owner page is open",
    "Invalid phone",
    "Enter invalid phone characters/format → Save",
    "Invalid phone is rejected or normalized according to the configured validation rules.")

# Mailing address
add("ADD-018","Different mailing address toggle OFF",
    "Add New Owner page is open",
    "Toggle OFF",
    "Keep Different mailing address OFF",
    "Mailing address fields remain hidden/disabled and assigned unit address remains the default mailing address behavior.")
add("ADD-019","Different mailing address toggle ON",
    "Add New Owner page is open",
    "Toggle ON",
    "Enable Different mailing address",
    "Additional mailing address inputs/section appear and can be completed.")
add("ADD-020","Different mailing address validation",
    "Different mailing address is enabled",
    "Incomplete address",
    "Enable toggle → leave required mailing address fields blank → Save",
    "Required mailing address validation appears if address fields are mandatory under the enabled option.")
add("ADD-021","Save owner with different mailing address",
    "Different mailing address is enabled and all required address data is valid",
    "Valid mailing address",
    "Enable toggle → enter valid address → assign unit → Save",
    "Owner is saved with the different mailing address associated correctly.")
add("ADD-022","Toggle mailing address OFF after entering values",
    "Different mailing address is ON and fields contain data",
    "Address values",
    "Enter address → toggle OFF → Save",
    "System follows the intended behavior: disabled address is not incorrectly stored/used as the owner's active mailing address.")

# Emergency contact
add("ADD-023","Emergency contact optional",
    "Add New Owner page is open",
    "No emergency contact",
    "Leave emergency contact blank → Save valid owner",
    "Owner can be created without an emergency contact.")
add("ADD-024","Emergency contact name only",
    "Add New Owner page is open",
    "Emergency first/last name; no phone/email",
    "Enter emergency contact name only → Save",
    "Validation prevents saving if a name requires at least phone or email, as indicated by the form.")
add("ADD-025","Emergency contact with phone",
    "Add New Owner page is open",
    "Emergency name + valid phone",
    "Enter name + phone → Save",
    "Emergency contact is saved successfully.")
add("ADD-026","Emergency contact with email",
    "Add New Owner page is open",
    "Emergency name + valid email",
    "Enter name + email → Save",
    "Emergency contact is saved successfully.")
add("ADD-027","Emergency contact with phone and email",
    "Add New Owner page is open",
    "Emergency name + phone + email",
    "Enter all emergency contact fields → Save",
    "Emergency contact is saved with all supplied information.")
add("ADD-028","Emergency contact invalid email",
    "Emergency contact is being added",
    "Invalid emergency email",
    "Enter emergency name + invalid email → Save",
    "Invalid email is rejected with validation.")
add("ADD-029","Emergency contact invalid phone",
    "Emergency contact is being added",
    "Invalid emergency phone",
    "Enter emergency name + invalid phone → Save",
    "Invalid phone is rejected/normalized according to phone validation rules.")

# Unit ownership
add("ADD-030","Property assignment required choice",
    "Add New Owner page is open",
    "Yes / No",
    "Review Unit ownership section",
    "User must choose Yes or No for assigning property.")
add("ADD-031","Select No for property assignment",
    "Add New Owner page is open",
    "No",
    "Select No",
    "Property/unit search and selection controls are hidden/disabled and owner can proceed without assignment if allowed.")
add("ADD-032","Select Yes for property assignment",
    "Add New Owner page is open",
    "Yes",
    "Select Yes",
    "Property/unit search and selectable unit list are displayed.")
add("ADD-033","Search unit by property name",
    "Yes is selected and units exist",
    "Five-Swpt-16",
    "Enter property name in property/unit search",
    "Matching units for the property are displayed.")
add("ADD-034","Search unit by unit name",
    "Yes is selected and units exist",
    "89",
    "Enter unit name in search",
    "Matching unit records are displayed.")
add("ADD-035","Search unit by address",
    "Yes is selected and units exist",
    "Streeterville, Chicago, Illinois, 60611",
    "Enter address in search",
    "Matching property/unit records are displayed.")
add("ADD-036","Search unit with no result",
    "Yes is selected",
    "Non-existing property/unit/address",
    "Enter non-existing search text",
    "No matching units are displayed and the empty state is clear.")
add("ADD-037","Select one unit",
    "Yes is selected and unit is available",
    "Five-Swpt-16 - 89",
    "Select unit checkbox",
    "Selected unit is visibly marked and will be associated with the owner upon successful save.")
add("ADD-038","Select multiple units",
    "Yes is selected and multiple units are available",
    "Unit 89, Unit 80, Unit 100",
    "Select multiple unit checkboxes",
    "All selected units are visibly selected and can be associated with the owner.")
add("ADD-039","Deselect unit",
    "A unit is selected",
    "Selected unit",
    "Click selected checkbox again",
    "Unit becomes deselected and will not be associated if not selected at save.")
add("ADD-040","Save with Yes but no unit selected",
    "Yes is selected",
    "No selected unit",
    "Select Yes → do not select any unit → Save",
    "System either blocks save with a clear validation message or follows the defined business rule; no unintended association is created.","High")
add("ADD-041","Prevent unauthorized property/unit selection",
    "User belongs to one HOA and another HOA has units",
    "Current HOA + other HOA units",
    "Open unit selector/search",
    "Only units belonging to the current authorized HOA are available.","Critical")
add("ADD-042","Prevent duplicate ownership association",
    "Owner/unit relationship already exists or business rule prevents duplicate association",
    "Existing owner + same unit",
    "Attempt to associate the same owner with an already-associated unit",
    "System prevents duplicate ownership association or applies the defined ownership rule without creating duplicate relationships.","Critical")

# Save/cancel
add("ADD-043","Save owner with valid mandatory data",
    "All required data is valid",
    "Unique owner + valid email + matching confirmation",
    "Fill required fields → choose property assignment option → Save owner",
    "Owner is created successfully and user receives a success confirmation.")
add("ADD-044","Cancel owner creation",
    "Add New Owner page is open",
    "Entered unsaved data",
    "Enter data → click Cancel",
    "User leaves the form without creating the owner; unsaved data is discarded according to the intended behavior.")
add("ADD-045","Cancel with unsaved changes confirmation",
    "Form contains unsaved data and product supports confirmation",
    "Unsaved data",
    "Enter data → click Cancel",
    "Confirmation appears if required by the product; user can remain or discard changes.")
add("ADD-046","Prevent double Save submission",
    "Valid form is ready to submit",
    "Valid owner",
    "Click Save owner repeatedly/quickly",
    "Only one owner is created; duplicate requests/records are not generated.","Critical")
add("ADD-047","Save failure handling",
    "Backend/API can return an error",
    "Valid form",
    "Submit during simulated server/API failure",
    "User sees a clear error, entered data is not silently lost, and no partial owner record is created.")
add("ADD-048","Form data retained after validation error",
    "Form contains multiple values",
    "Invalid email or mismatch",
    "Submit invalid form",
    "Validation is shown while valid entered fields remain populated.")

# Backend / invitation / account status
add("ADD-049","Owner record created in database",
    "Valid owner submission",
    "Unique email + owner details",
    "Create owner → verify owner list/database through authorized QA tooling",
    "Exactly one owner record is created with correct owner information.","Critical")
add("ADD-050","Owner invitation is generated",
    "Owner is successfully created with an email",
    "Valid unique owner email",
    "Create owner → check invitation workflow",
    "Owner receives/is queued for the intended invitation and account status follows the configured invitation lifecycle.")
add("ADD-051","New owner account status - Invitation Sent",
    "Invitation has been generated/sent",
    "New owner",
    "Create owner → open Property Owners list",
    "Owner shows **Invitation Sent** when the invitation has been sent and no verification has occurred.")
add("ADD-052","Invitation acceptance changes status",
    "Owner has received invitation",
    "Invitation link",
    "Open invitation → complete owner account activation/verification",
    "Owner status changes according to the configured lifecycle, expected to reach **Verified** after successful verification.")
add("ADD-053","Pending Verification status",
    "Owner exists in the applicable verification state",
    "Pending verification owner",
    "Create/trigger the state → open owner list",
    "Owner shows **Pending Verification** when verification is required but not completed, according to the configured lifecycle.")
add("ADD-054","Verified status",
    "Owner has completed the required verification",
    "Verified owner",
    "Complete verification → open owner list",
    "Owner shows **Verified**.")
add("ADD-055","Invitation resend behavior",
    "Owner is in Invitation Sent/Pending Verification state and resend is supported",
    "Existing invited owner",
    "Trigger resend invitation",
    "Invitation is resent according to the configured workflow without creating a duplicate owner account.")
add("ADD-056","Invitation email uses correct recipient",
    "Owner creation succeeds",
    "Unique owner email",
    "Create owner → inspect invitation delivery target using test mail tooling",
    "Invitation is addressed to the owner's stored email, not an emergency contact or another address.","Critical")

# Invoice automation / Manage HOA settings
add("ADD-057","Assign owner to unit triggers invoice creation according to Manage HOA settings",
    "Manage HOA Fee settings are configured for invoice generation; valid owner/unit assignment",
    "Owner + assigned unit + configured fee schedule",
    "Create owner → select Yes → assign unit → Save → inspect invoices",
    "Invoices are created according to the configured Manage HOA Fee settings, schedule, amount, due date, and applicable rules.","Critical")
add("ADD-058","No invoice when invoice generation is not applicable",
    "Manage HOA Fee settings are configured so no invoice should be generated for the tested condition",
    "Owner + unit + fee configuration",
    "Create/assign owner to unit → inspect invoices",
    "No unexpected invoice is generated when configuration/rules do not call for one.","Critical")
add("ADD-059","Verify invoice belongs to correct owner/unit",
    "Invoice is generated after owner/unit assignment",
    "Owner + assigned unit",
    "Create owner → assign unit → inspect generated invoice",
    "Generated invoice is linked to the correct HOA, owner, property/unit, amount, and configured billing period.","Critical")
add("ADD-060","Verify no duplicate invoice creation on repeated save/retry",
    "Owner creation/invoice workflow is active",
    "Same owner/unit and billing period",
    "Submit once and simulate retry/refresh as applicable → inspect invoices",
    "Only the expected invoice is created; duplicate invoice records are not generated.","Critical")
add("ADD-061","Verify invoice date/due date follows Manage HOA settings",
    "Manage HOA Fee has defined billing/due-date settings",
    "Configured invoice schedule",
    "Create owner + assign unit → inspect generated invoice dates",
    "Invoice issue/due dates follow configured Manage HOA settings and applicable timezone/date rules.","Critical")
add("ADD-062","Verify invoice amount follows Manage HOA settings",
    "Manage HOA Fee has a configured fee amount",
    "Configured fee amount",
    "Create owner + assign unit → inspect generated invoice",
    "Invoice amount matches the applicable Manage HOA Fee configuration.","Critical")

# Post creation list / filter integration
add("ADD-063","New owner appears in Property Owners list",
    "Owner was created successfully",
    "New owner",
    "Return to Property Owners → refresh",
    "New owner appears with correct name, property/unit, effective date, contact, and account status.")
add("ADD-064","New owner searchable by owner name",
    "Owner was created successfully",
    "New owner name",
    "Search new owner by name",
    "New owner is returned.")
add("ADD-065","New owner searchable by property/unit",
    "Owner was created and assigned to a unit",
    "Assigned property/unit",
    "Use Property Owners filter/search for the assigned property/unit",
    "New owner appears in the matching results.")
add("ADD-066","New owner appears under correct account status filter",
    "Owner was created and invitation/verification state is known",
    "Invitation Sent / Pending Verification / Verified",
    "Apply the matching status filter",
    "New owner appears under the correct status and not under unrelated statuses.")

# Data integrity / permissions
add("ADD-067","Required field validation does not create partial owner",
    "One or more required fields are invalid/missing",
    "Invalid submission",
    "Attempt Save → inspect owner list/data",
    "No partial/incomplete owner record is created.")
add("ADD-068","Board Member permission to create owner",
    "Logged in as authorized HOA Board Member",
    "Board Member role",
    "Open Add Owner and submit valid data",
    "Authorized Board Member can create the owner.")
add("ADD-069","Unauthorized role cannot create owner",
    "Test account lacks owner-create permission",
    "Unauthorized role",
    "Attempt to access Add Owner directly and/or submit",
    "Access/create operation is denied according to role permissions.","Critical")
add("ADD-070","Cross-HOA owner creation is prevented",
    "User is scoped to one HOA",
    "Another HOA property/unit",
    "Attempt to assign owner to property/unit outside current HOA",
    "System prevents cross-HOA assignment and no unauthorized relationship is created.","Critical")

# UI/UX
add("ADD-071","Verify Cancel and Save owner buttons",
    "Add New Owner page is open",
    "Cancel / Save owner",
    "Review bottom action area",
    "Both actions are visible, enabled/disabled appropriately, and behave correctly.")
add("ADD-072","Verify required field indicators",
    "Add New Owner page is open",
    "Required fields",
    "Review labels",
    "Required fields show a clear required indicator.")
add("ADD-073","Verify help text for email confirmation",
    "Add New Owner page is open",
    "Email fields",
    "Review Owner information section",
    "Instruction to re-type email to avoid invite typos is visible/consistent with the design.")
add("ADD-074","Verify international phone country selector",
    "Add New Owner page is open",
    "US and supported country codes",
    "Open phone country selector",
    "Supported country codes can be selected and phone formatting follows the selected country.")
add("ADD-075","Verify unit search placeholder",
    "Yes is selected",
    "Property/unit/address search",
    "Review search field",
    "Search field indicates property name, unit name, or address can be searched.")
add("ADD-076","Verify selected unit remains selected after search",
    "Yes selected and at least one unit selected",
    "Selected Unit 89",
    "Select unit → search another term → clear search",
    "Previously selected unit remains selected unless product explicitly resets selections.")
add("ADD-077","Verify selected units remain selected after scrolling",
    "Multiple units are available",
    "Multiple selected units",
    "Select units → scroll through list → return",
    "Selected state remains correct.")
add("ADD-078","Verify form does not submit on Enter unintentionally",
    "Form contains incomplete/partially completed data",
    "Keyboard Enter",
    "Press Enter in different fields",
    "Form follows intended keyboard behavior and does not create an incomplete owner unexpectedly.")

# End-to-end
add("ADD-079","End-to-end create owner → assign unit → invoice → invitation → status",
    "Manage HOA Fee configured; test email available; valid Board Member",
    "Unique owner + unique email + selected unit",
    "Enter owner information → confirm email → optional phone/emergency contact → select Yes → search/select unit → Save → verify owner list → verify invoice → verify invitation → verify initial account status → complete invitation/verification if testable",
    "Exactly one owner is created, selected unit is associated, invoice generation follows Manage HOA settings, invitation is generated to the correct email, and account status progresses through the configured lifecycle (Invitation Sent / Pending Verification / Verified).","Critical")

# Write sheet
for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c, h)
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin = Side(style="thin", color="D9E1F2")
for r, row in enumerate(cases, 2):
    for c, value in enumerate(row, 1):
        cell = ws.cell(r, c, value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[r].height = 58

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:N{ws.max_row}"

widths = {
    "A":12, "B":22, "C":42, "D":36, "E":32, "F":58,
    "G":72, "H":42, "I":16, "J":14, "K":14, "L":18, "M":26, "N":34
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

status_dv = DataValidation(type="list", formula1='"Not Started,Pass,Fail,Blocked,Not Applicable,Retest"', allow_blank=False)
priority_dv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=False)
ws.add_data_validation(status_dv)
ws.add_data_validation(priority_dv)
status_dv.add(f"I2:I{ws.max_row}")
priority_dv.add(f"J2:J{ws.max_row}")

ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Pass"'], fill=PatternFill("solid", fgColor="C6EFCE"))
)
ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Fail"'], fill=PatternFill("solid", fgColor="FFC7CE"))
)
ws.conditional_formatting.add(
    f"I2:I{ws.max_row}",
    FormulaRule(formula=['I2="Blocked"'], fill=PatternFill("solid", fgColor="FFEB9C"))
)
ws.conditional_formatting.add(
    f"J2:J{ws.max_row}",
    FormulaRule(formula=['J2="Critical"'], fill=PatternFill("solid", fgColor="F4CCCC"))
)
ws.sheet_view.showGridLines = False

# Dedicated summary
if "Add Owner Summary" in wb.sheetnames:
    del wb["Add Owner Summary"]
sm = wb.create_sheet("Add Owner Summary")
sm["A1"] = "Add New Owner - QA Summary"
sm["A1"].font = Font(size=18, bold=True)
sm.merge_cells("A1:D1")

summary = [
    ("Total Test Cases", f"=COUNTA('Add Owner QA'!A2:A{ws.max_row})"),
    ("Passed", f'=COUNTIF(\'Add Owner QA\'!I2:I{ws.max_row},"Pass")'),
    ("Failed", f'=COUNTIF(\'Add Owner QA\'!I2:I{ws.max_row},"Fail")'),
    ("Blocked", f'=COUNTIF(\'Add Owner QA\'!I2:I{ws.max_row},"Blocked")'),
    ("Not Started", f'=COUNTIF(\'Add Owner QA\'!I2:I{ws.max_row},"Not Started")'),
    ("Pass Rate", '=IF(B2=0,0,B3/B2)'),
    ("Execution Rate", '=IF(B2=0,0,(B3+B4+B5)/B2)'),
    ("Critical Cases", f'=COUNTIF(\'Add Owner QA\'!J2:J{ws.max_row},"Critical")'),
    ("Critical Failures", f'=COUNTIFS(\'Add Owner QA\'!J2:J{ws.max_row},"Critical",\'Add Owner QA\'!I2:I{ws.max_row},"Fail")'),
]
for i, (label, formula) in enumerate(summary, 2):
    sm.cell(i,1,label).font = Font(bold=True)
    sm.cell(i,2,formula)
    sm.cell(i,1).fill = PatternFill("solid", fgColor="D9EAF7")
    sm.cell(i,1).border = Border(bottom=thin)
    sm.cell(i,2).border = Border(bottom=thin)

sm["B7"].number_format = "0.0%"
sm["B8"].number_format = "0.0%"
sm.column_dimensions["A"].width = 30
sm.column_dimensions["B"].width = 18
sm.sheet_view.showGridLines = False

# Add important workflow notes
if "Test Reference" in wb.sheetnames:
    ref = wb["Test Reference"]
else:
    ref = wb.create_sheet("Test Reference")

next_row = ref.max_row + 2
ref.cell(next_row, 1, "Add New Owner workflow")
ref.cell(next_row, 1).font = Font(bold=True)
ref.cell(next_row + 1, 1, "Owner creation")
ref.cell(next_row + 1, 2, "First Name + Last Name + unique Email + Confirm Email; Phone is optional.")
ref.cell(next_row + 2, 1, "Unit ownership")
ref.cell(next_row + 2, 2, "Yes → search property/unit/address → select one or more units; No → no unit assignment.")
ref.cell(next_row + 3, 1, "Invoice automation")
ref.cell(next_row + 3, 2, "If owner is assigned to a unit, verify invoice creation follows Manage HOA Fee settings.")
ref.cell(next_row + 4, 1, "Invitation")
ref.cell(next_row + 4, 2, "Verify owner invitation is generated/sent to the stored owner email.")
ref.cell(next_row + 5, 1, "Account status")
ref.cell(next_row + 5, 2, "Validate Invitation Sent, Pending Verification, and Verified according to the configured lifecycle.")
for r in range(next_row, next_row + 6):
    ref.cell(r, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    ref.cell(r, 1).font = Font(bold=True)
    ref.cell(r, 2).alignment = Alignment(wrap_text=True)

# Update main summary with Add Owner count
dash = wb["Summary Dashboard"]
row = dash.max_row + 2
dash.cell(row,1,"Add New Owner QA")
dash.cell(row,1).font = Font(bold=True)
dash.cell(row,2,f"='Add Owner Summary'!B2")
dash.cell(row,3,"test cases added")
dash.cell(row,1).fill = PatternFill("solid", fgColor="D9EAF7")

wb.save(path)

print(f"Updated: {path}")
print(f"Add New Owner test cases added: {len(cases)}")
