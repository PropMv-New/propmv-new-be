from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from pathlib import Path

src = Path("/mnt/data/HOA_Dashboard_QA_Tracker_v3.xlsx")
out = Path("/mnt/data/HOA_Dashboard_QA_Tracker_v4.xlsx")
wb = load_workbook(src)

if "Add Property QA" in wb.sheetnames:
    del wb["Add Property QA"]
if "Add Property Summary" in wb.sheetnames:
    del wb["Add Property Summary"]

ws = wb.create_sheet("Add Property QA")
headers = ["TC ID","Module","Scenario","Preconditions","Test Data","Steps","Expected Result","Actual Result","Status","Priority","Bug ID","Tester","Evidence","Remarks"]
ws.append(headers)

cases = []
def add(tc, module, scenario, pre, data, steps, expected, priority="P0"):
    cases.append([tc,module,scenario,pre,data,steps,expected,"","Not Tested",priority,"","","",""])

# Mandatory fields / navigation
add("PROP-ADD-001","Add Property","Open Add New Property","User has permission to create properties","N/A","Open Properties → Add New Property","Add New Property page opens correctly.")
add("PROP-ADD-002","Mandatory","Property name is mandatory","Add Property page open","Blank property name","Leave Property Name blank → submit","Validation is displayed and property is not created.")
add("PROP-ADD-003","Mandatory","Unit number is mandatory","Add Property page open","Blank unit number","Leave Unit Number blank → submit","Validation is displayed and property is not created.")
add("PROP-ADD-004","Mandatory","Address is mandatory","Add Property page open","Blank address fields","Leave required address fields blank → submit","Address validation is displayed and property is not created.")
add("PROP-ADD-005","Mandatory","Deposit account is mandatory","Add Property page open; HOA bank account exists","No deposit account selected","Do not select deposit account → submit","Validation is displayed and property is not created.")
add("PROP-ADD-006","Mandatory","All mandatory fields completed","Add Property page open; valid bank account exists","Valid property/unit/address/bank account","Fill all mandatory fields → save","Property is created successfully.")
add("PROP-ADD-007","Validation","Property name duplicate across communities","Existing same property name in another community","Same property name","Create in different community","System follows stated rule that property names may repeat across communities.")
add("PROP-ADD-008","Validation","Unit name duplicate within same property","Property already has same unit","Existing property + same unit","Create duplicate unit in same property","System prevents duplicate unit name within the same property.")
add("PROP-ADD-009","Validation","Unique unit within property","Property exists","New unit number","Create property/unit","Unit is accepted when unique within the property.")
add("PROP-ADD-010","Validation","Save with only mandatory fields","Valid mandatory values","No optional data","Save","Property is created without requiring optional fields.")
add("PROP-ADD-011","Navigation","Cancel Add Property","Add Property page open","Unsaved changes","Enter data → Cancel","User returns without creating the property; unsaved data is not persisted.")
add("PROP-ADD-012","Navigation","Progress indicator","Add Property page open","N/A","Move through form sections","Progress indicator updates consistently with completed required sections.")

# Address / Google Places
add("PROP-ADDR-001","Google Places","Address autocomplete","Google Places API available","Valid US address","Start typing address","Relevant Google Places suggestions are displayed.")
add("PROP-ADDR-002","Google Places","Select suggested place","Suggestions displayed","Valid place suggestion","Select a suggestion","Address fields are populated with the selected place data.")
add("PROP-ADDR-003","Google Places","City populated from selected place","Valid place selected","City","Select place","City is populated correctly.")
add("PROP-ADDR-004","Google Places","State populated from selected place","Valid place selected","State","Select place","State is populated correctly.")
add("PROP-ADDR-005","Google Places","ZIP populated from selected place","Valid place selected","ZIP","Select place","ZIP code is populated correctly.")
add("PROP-ADDR-006","Google Places","Address Line 1 populated","Valid place selected","Street address","Select place","Address Line 1 is populated correctly.")
add("PROP-ADDR-007","Google Places","Address Line 2 optional","Valid place selected","Apt/Unit/Suite","Enter Address Line 2","Optional address line is accepted.")
add("PROP-ADDR-008","Google Places","Manual/invalid address validation","Address entry available","Invalid/incomplete address","Enter invalid address → submit","System validates required address data and does not save invalid/incomplete required address.")
add("PROP-ADDR-009","Google Places","API unavailable","Google Places unavailable/timeout","N/A","Attempt address lookup","User receives a clear error/fallback behavior; form does not silently save incorrect address.")
add("PROP-ADDR-010","Google Places","Address persistence","Valid address selected","Known address","Select place → save property → reopen","Saved address matches selected address.")

# Financial setup
add("PROP-FIN-001","Starting Balance","Starting Balance required when owner selected","Owner is selected","Blank starting balance","Select owner → leave starting balance blank → save","Starting Balance validation is displayed; property is not saved until required value is supplied.")
add("PROP-FIN-002","Starting Balance","Balance as of Date required when owner selected","Owner is selected","Blank date","Select owner → enter balance → leave date blank → save","Balance as of Date validation is displayed.")
add("PROP-FIN-003","Starting Balance","Positive starting balance","Owner selected","$500","Enter positive balance + date → save","Balance is accepted and stored correctly.")
add("PROP-FIN-004","Starting Balance","Zero starting balance","Owner selected","$0","Enter zero + date → save","Zero balance is accepted according to the form rules.")
add("PROP-FIN-005","Starting Balance","Negative starting balance","Owner selected","-$500","Enter negative balance + date → save","Negative value is handled according to the stated starting-balance rule (advance/credit) and stored correctly.")
add("PROP-FIN-006","Starting Balance","Invalid starting balance format","Owner selected","Letters/symbols","Enter invalid amount → save","Validation prevents invalid amount.")
add("PROP-FIN-007","Starting Balance","Balance as of Date valid date","Owner selected","Valid historical date","Enter balance + date → save","Date is accepted and persisted.")
add("PROP-FIN-008","Starting Balance","Future Balance as of Date","Owner selected","Future date","Enter balance + future date → save","System follows defined business rule for future dates; invalid dates are rejected if prohibited.")
add("PROP-FIN-009","Starting Balance","Starting balance persistence","Owner selected; property created","Known amount/date","Create property → reopen","Starting balance and balance date are retained accurately.")
add("PROP-FIN-010","Starting Balance","No owner selected","No owner assigned","Starting balance/date blank","Create property without owner","Starting Balance section is not incorrectly made mandatory when no owner is assigned.")

# Deposit account
add("PROP-BANK-001","Deposit Account","Bank account list loads","At least one HOA bank account exists","Valid account","Open Deposit Bank Account section","Available accounts are displayed.")
add("PROP-BANK-002","Deposit Account","Select deposit account","Bank account exists","Account A","Select account → save","Selected deposit account is associated with the property/unit.")
add("PROP-BANK-003","Deposit Account","No account available","No HOA bank account configured","N/A","Open Add Property","System clearly indicates deposit account is required/unavailable and does not allow invalid creation.")
add("PROP-BANK-004","Deposit Account","Selected account persistence","Account selected","Account A","Save → reopen property","Same deposit account is displayed.")
add("PROP-BANK-005","Deposit Account","Correct account for owner payments","Account selected","Account A","Create property → record owner payment","Payment routing uses the selected deposit account as configured.")

# Owner association
add("PROP-OWN-001","Owner Association","Search owner by name","Owners exist","Owner full/partial name","Type in Search Owners","Matching owners are displayed.")
add("PROP-OWN-002","Owner Association","Search owner by email","Owners exist","Owner email/full/partial","Search by email","Matching owner is displayed.")
add("PROP-OWN-003","Owner Association","Search no results","Owners exist","Invalid name/email","Search","Correct empty state is displayed.")
add("PROP-OWN-004","Owner Association","Select existing verified owner","Verified owner exists","Verified owner","Search → select owner → save","Owner is associated and billing workflow is eligible for invoice generation.")
add("PROP-OWN-005","Owner Association","Select pending/unverified owner","Unverified owner exists","Pending owner","Search → select owner → save","Owner can be associated if allowed, but invoice is not generated until owner is verified.")
add("PROP-OWN-006","Owner Association","Multiple owners on one unit","Multiple owners exist","Owner A + Owner B","Select multiple owners → save","Multiple owners are associated with the same unit without overwriting each other.")
add("PROP-OWN-007","Owner Association","Remove one of multiple owners","Multiple owners associated","Owner A + B","Remove Owner A → save","Owner B remains associated; Owner A is removed.")
add("PROP-OWN-008","Owner Association","Create new owner from Add Property","Create owner permission available","Valid owner data","Click Create a new owner → complete form → create","New owner is created and becomes selectable/associated according to workflow.")
add("PROP-OWN-009","Owner Association","Owner search after creation","New owner created","New owner name/email","Search Owners","New owner appears in search results.")
add("PROP-OWN-010","Owner Association","Owner association persistence","Owner selected","Known owner","Save → reopen property","Correct owner association persists.")

# Create owner modal / inline
add("PROP-OWNER-001","Create Owner","First Name required","Create New Owner modal open","Blank first name","Submit blank first name","Validation displayed.")
add("PROP-OWNER-002","Create Owner","Last Name required","Create New Owner modal open","Blank last name","Submit blank last name","Validation displayed.")
add("PROP-OWNER-003","Create Owner","Email required","Create New Owner modal open","Blank email","Submit blank email","Validation displayed.")
add("PROP-OWNER-004","Create Owner","Email confirmation required","Create New Owner modal open","Blank confirmation","Enter email only → create","Validation displayed.")
add("PROP-OWNER-005","Create Owner","Email confirmation mismatch","Create New Owner modal open","Two different emails","Enter different email/confirmation → create","Mismatch validation displayed; owner is not created.")
add("PROP-OWNER-006","Create Owner","Unique owner email","Owner email already exists","Existing email","Create owner with existing email","System prevents duplicate owner email according to unique-email rule.")
add("PROP-OWNER-007","Create Owner","Phone duplicate behavior","Owner phone already exists","Existing phone","Create another owner with same phone","System follows phone duplicate validation rule and displays appropriate result.")
add("PROP-OWNER-008","Create Owner","Emergency contact optional","Create owner modal open","No emergency contact","Fill required owner data → create","Owner can be created without emergency contact.")
add("PROP-OWNER-009","Create Owner","Emergency contact with name requires phone or email","Create owner modal open","Emergency name + no phone/email","Enter emergency name only → create","Validation requires at least a phone or email when a name is supplied.")
add("PROP-OWNER-010","Create Owner","Mailing address toggle","Create owner modal open","Toggle ON","Enable mailing address different from unit","Additional mailing-address fields/flow are displayed and saved correctly.")
add("PROP-OWNER-011","Create Owner","Cancel create owner","Create owner modal open","Unsaved owner data","Enter data → Cancel","Owner is not created.")
add("PROP-OWNER-012","Create Owner","Create owner success","Valid required owner data","Valid unique email","Complete form → Create Owner","Owner is created successfully and available for association.")

# Auto invoice generation
add("PROP-INV-001","Invoice Generation","Verified owner triggers invoice","Manage HOA Fee configured; verified owner selected","Starting balance/date + fee config","Create property with verified owner","Invoice is auto-generated according to Manage HOA Fee configuration.")
add("PROP-INV-002","Invoice Generation","Unverified owner does not trigger invoice","Manage HOA Fee configured; owner unverified","Pending/invitation owner","Create property with unverified owner","No invoice is generated until owner becomes verified.")
add("PROP-INV-003","Invoice Generation","No owner means no owner-triggered invoice","Manage HOA Fee configured","No owner","Create property without owner","No owner-triggered invoice is generated.")
add("PROP-INV-004","Invoice Generation","Invoice amount uses Manage HOA Fee","Manage HOA Fee configured","Fee $100/month; owner verified","Create property with verified owner","Generated invoice amount follows configured HOA fee.")
add("PROP-INV-005","Invoice Generation","Invoice schedule uses Manage HOA Fee due day","Manage HOA Fee configured","Due day 13th","Create property with verified owner","Generated invoice due date follows configured due-day rules.")
add("PROP-INV-006","Invoice Generation","Starting balance/date affects invoice generation","Manage HOA Fee configured; verified owner","Balance $500; date specified","Create property with owner","Invoice generation reflects starting balance/date according to billing rules.")
add("PROP-INV-007","Invoice Generation","Multiple owners invoice behavior","Manage HOA Fee configured; multiple verified owners","Two verified owners on one unit","Create property/unit with both owners","Invoice generation follows the defined ownership/billing rule without duplicate or unintended invoices.")
add("PROP-INV-008","Invoice Generation","Verified after initial creation","Property created with unverified owner","Owner initially pending → later verified","Verify owner → inspect invoices","Invoice is generated after verification according to the configured workflow, if verification is the triggering event.")
add("PROP-INV-009","Invoice Generation","No duplicate invoice on repeated save","Verified owner + configured fee","Same property save/retry","Save/retry workflow multiple times","System does not create duplicate invoices for the same billing period/event.")
add("PROP-INV-010","Invoice Generation","Invoice data persistence","Invoice auto-generated","Known invoice","Create property → inspect invoice","Invoice is linked to correct owner/unit/property and amount/date are correct.")

# Tags
add("PROP-TAG-001","Tags","Add tag","Add Property page open","Corner Unit","Enter tag → add","Tag is added to the unit/property.")
add("PROP-TAG-002","Tags","Multiple tags","Add Property page open","Corner Unit, Pool View","Add multiple tags","All selected tags are displayed and persisted.")
add("PROP-TAG-003","Tags","Duplicate tag","Existing tag on same unit","Same tag twice","Add same tag twice","System prevents duplicate tag assignment or handles it consistently.")
add("PROP-TAG-004","Tags","Remove tag","Tag exists","Existing tag","Remove tag → save","Tag is removed.")
add("PROP-TAG-005","Tags","Tag persistence","Tags added","Known tags","Save → reopen","Tags remain associated.")
add("PROP-TAG-006","Tags","Tag filtering/reporting","Tagged units exist","Known tag","Use property/unit filter/reporting by tag","Tagged units are returned correctly.")

# Integration / regression
add("PROP-E2E-001","End-to-End","Create property with all mandatory fields and verified owner","Bank account, Manage HOA Fee, verified owner available","Property + unit + address + bank + balance/date + owner","Complete form → save","Property/unit created, owner associated, and invoice generated according to configuration.")
add("PROP-E2E-002","End-to-End","Create property with mandatory fields but no owner","Bank account available","Mandatory data only","Complete required fields → save","Property created; no owner-triggered invoice.")
add("PROP-E2E-003","End-to-End","Create property with unverified owner","Bank + Manage HOA Fee configured; unverified owner exists","All required data + pending owner","Save property","Property/owner association succeeds if allowed; invoice is not generated.")
add("PROP-E2E-004","End-to-End","Create property with multiple verified owners","Bank + fee config + verified owners","Owner A + Owner B","Assign both → save","Both owners remain associated and invoice behavior follows defined multi-owner rule.")
add("PROP-E2E-005","End-to-End","Create owner from property then associate","Bank account + fee config","New owner","Create new owner → select owner → save property","Owner is created, associated, and billing behavior follows verification state.")
add("PROP-E2E-006","End-to-End","Google Place address + verified owner + invoice","Google Places available; fee configured","Valid place + verified owner","Select address → assign owner → save","Correct address saved and invoice generated according to configuration.")
add("PROP-E2E-007","End-to-End","Starting balance and date regression","Fee configured; verified owner","Known balance/date","Create property → inspect generated invoice","Starting balance/date are reflected correctly and no duplicate invoice occurs.")
add("PROP-E2E-008","End-to-End","Property creation blocked without deposit account","No deposit account available","Mandatory property/address/unit, no bank","Attempt save","Creation is blocked with clear deposit-account validation.")
add("PROP-E2E-009","End-to-End","Property creation after Manage HOA Fee completion","Manage HOA Fee initially incomplete then completed","Valid fee config","Attempt create before completion → complete fee setup → retry","Dependent workflow is blocked before completion and available after completion.")
add("PROP-E2E-010","End-to-End","Property data appears in Properties listing","Property created","Known property","Create → return to listing","New property/unit appears with correct address, owner and balances.")
add("PROP-E2E-011","End-to-End","Search newly created property","Property created","Property/unit/owner name","Use listing search","New record is searchable by supported fields.")
add("PROP-E2E-012","End-to-End","Refresh after creation","Property created","N/A","Create → refresh listing","Record remains visible with current values.")

for row in cases:
    ws.append(row)

navy="1F4E78"; white="FFFFFF"; thin=Side(style="thin",color="B7B7B7")
for cell in ws[1]:
    cell.fill=PatternFill("solid",fgColor=navy)
    cell.font=Font(color=white,bold=True)
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    cell.border=Border(bottom=thin)

for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
    for cell in row:
        cell.alignment=Alignment(vertical="top",wrap_text=True)
        cell.border=Border(bottom=thin)

ws.freeze_panes="A2"
ws.auto_filter.ref=f"A1:N{ws.max_row}"
ws.sheet_view.showGridLines=False
for i,w in enumerate([18,20,38,42,34,48,62,42,16,12,14,18,24,34],1):
    ws.column_dimensions[get_column_letter(i)].width=w

status=DataValidation(type="list",formula1='"Not Tested,In Progress,Passed,Failed,Blocked,Not Applicable,Retest Required"',allow_blank=False)
priority=DataValidation(type="list",formula1='"P0,P1,P2"',allow_blank=False)
ws.add_data_validation(status); status.add(f"I2:I{ws.max_row}")
ws.add_data_validation(priority); priority.add(f"J2:J{ws.max_row}")

for s,c in {"Passed":"E2F0D9","Failed":"F4CCCC","Blocked":"FCE4D6","In Progress":"FFF2CC","Not Tested":"E7E6E6","Retest Required":"FCE4D6","Not Applicable":"D9EAF7"}.items():
    ws.conditional_formatting.add(f"I2:I{ws.max_row}",FormulaRule(formula=[f'$I2="{s}"'],fill=PatternFill("solid",fgColor=c)))
for p,c in {"P0":"F4CCCC","P1":"FFF2CC","P2":"E2F0D9"}.items():
    ws.conditional_formatting.add(f"J2:J{ws.max_row}",FormulaRule(formula=[f'$J2="{p}"'],fill=PatternFill("solid",fgColor=c)))

sm=wb.create_sheet("Add Property Summary")
sm.sheet_view.showGridLines=False
sm["A1"]="Add Property QA Summary"
sm["A1"].font=Font(size=18,bold=True)
sm.merge_cells("A1:D1")
summary=[
("Total Test Cases",f"=COUNTA('Add Property QA'!A2:A{ws.max_row})"),
("Passed",f'=COUNTIF(\'Add Property QA\'!I2:I{ws.max_row},"Passed")'),
("Failed",f'=COUNTIF(\'Add Property QA\'!I2:I{ws.max_row},"Failed")'),
("Blocked",f'=COUNTIF(\'Add Property QA\'!I2:I{ws.max_row},"Blocked")'),
("Not Tested",f'=COUNTIF(\'Add Property QA\'!I2:I{ws.max_row},"Not Tested")'),
("Coverage %","=IF(B2=0,0,(B3+B4+B5)/B2)"),
("Pass Rate %","=IF((B3+B4+B5)=0,0,B3/(B3+B4+B5))"),
("P0 Failed/Blocked",f'=COUNTIFS(\'Add Property QA\'!J2:J{ws.max_row},"P0",\'Add Property QA\'!I2:I{ws.max_row},"Failed")+COUNTIFS(\'Add Property QA\'!J2:J{ws.max_row},"P0",\'Add Property QA\'!I2:I{ws.max_row},"Blocked")')
]
for i,(label,formula) in enumerate(summary,2):
    sm.cell(i,1,label).font=Font(bold=True)
    sm.cell(i,2,formula)
    sm.cell(i,1).fill=PatternFill("solid",fgColor="D9EAF7")
sm["B7"].number_format="0.00%"; sm["B8"].number_format="0.00%"
sm.column_dimensions["A"].width=28; sm.column_dimensions["B"].width=20

# Add a short source/requirements sheet entry
ref = wb["Test Reference"] if "Test Reference" in wb.sheetnames else wb.create_sheet("Test Reference")
r=ref.max_row+2
ref.cell(r,1,"Add Property QA coverage")
ref.cell(r,1).font=Font(bold=True)
ref.cell(r,2,"Covers mandatory property/unit/address/deposit fields, Google Places address flow, starting balance/date when an owner is selected, owner search and creation, multiple owners, tags, verification-dependent invoice generation, Manage HOA Fee integration, and regression/end-to-end workflows.")
ref.cell(r,2).alignment=Alignment(wrap_text=True)

wb.save(out)
print(f"Added {len(cases)} Add Property test cases.")
print(out)
